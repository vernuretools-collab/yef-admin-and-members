"""Generate the YEF member guide PowerPoint."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1B, 0x2E, 0x6B)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x5C, 0x60, 0x7A)
INK = RGBColor(0x1A, 0x1D, 0x2E)
LIGHT = RGBColor(0xEE, 0xF0, 0xF7)
PALE = RGBColor(0xDC, 0xE1, 0xF5)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run(run, text, size=18, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_round(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, l, t, w, h, items, size=18, color=INK, spacing=10):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        p.level = 0
        set_run(p.add_run(), "•  " + item, size=size, bold=False, color=color)
    return box


def header_bar(slide, title, subtitle=None):
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.15), NAVY)
    add_rect(slide, 0, Inches(1.15), SLIDE_W, Inches(0.08), GOLD)
    add_textbox(slide, Inches(0.6), Inches(0.22), Inches(11.5), Inches(0.5), title, 28, True, WHITE)
    if subtitle:
        add_textbox(
            slide, Inches(0.6), Inches(0.7), Inches(11.5), Inches(0.35), subtitle, 14, False, PALE
        )


def footer(slide, n, total=12):
    add_rect(slide, 0, Inches(7.22), SLIDE_W, Inches(0.28), NAVY)
    add_textbox(
        slide,
        Inches(0.5),
        Inches(7.22),
        Inches(9),
        Inches(0.28),
        "YEF Network  ·  Member guide",
        11,
        False,
        WHITE,
    )
    add_textbox(
        slide,
        Inches(10.5),
        Inches(7.22),
        Inches(2.4),
        Inches(0.28),
        f"{n}  /  {total}",
        11,
        False,
        WHITE,
        PP_ALIGN.RIGHT,
    )


def card(slide, l, t, w, h, title, body, title_size=16, body_size=14):
    add_round(slide, l, t, w, h, LIGHT)
    add_textbox(slide, l + Inches(0.22), t + Inches(0.16), w - Inches(0.4), Inches(0.4), title, title_size, True, NAVY)
    add_textbox(slide, l + Inches(0.22), t + Inches(0.55), w - Inches(0.4), h - Inches(0.7), body, body_size, False, SLATE)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1 Title
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    add_rect(s, 0, 0, Inches(0.22), SLIDE_H, GOLD)
    add_textbox(s, Inches(0.9), Inches(1.7), Inches(11), Inches(0.4), "YEF NETWORK", 16, True, GOLD)
    add_textbox(
        s,
        Inches(0.9),
        Inches(2.15),
        Inches(11.5),
        Inches(1.6),
        "How to log in and use\nReferrals, Thank You, and One-to-Ones",
        36,
        True,
        WHITE,
    )
    add_textbox(
        s,
        Inches(0.9),
        Inches(4.4),
        Inches(11),
        Inches(0.8),
        "A short member guide. Follow the steps in the app — no admin screens.",
        18,
        False,
        PALE,
    )
    add_textbox(
        s,
        Inches(0.9),
        Inches(6.5),
        Inches(11),
        Inches(0.4),
        "Invitation-only  ·  Member dashboard",
        14,
        False,
        GOLD,
    )

    # 2 Login
    s = prs.slides.add_slide(blank)
    header_bar(s, "How to log in", "Access is by invitation only")
    footer(s, 2)
    add_bullets(
        s,
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5.3),
        [
            "Your chapter admin sends you login credentials. There is no self-signup.",
            "Open the YEF login page (the /login screen).",
            "Enter your email address and password.",
            "Click Sign in.",
            "Members go straight to the member dashboard.",
            "If you have not received credentials, contact your chapter admin.",
        ],
        20,
        INK,
        14,
    )

    # 3 After login
    s = prs.slides.add_slide(blank)
    header_bar(s, "After you log in", "Use the left sidebar to move around")
    footer(s, 3)
    card(
        s,
        Inches(0.55),
        Inches(1.55),
        Inches(3.9),
        Inches(4.9),
        "Dashboard",
        "Home for My Activity. This is where you log referrals, thank-yous (TYFCBs), and one-to-ones.",
        20,
        16,
    )
    card(
        s,
        Inches(4.7),
        Inches(1.55),
        Inches(3.9),
        Inches(4.9),
        "My Referrals",
        "See slips you gave or received. Convert a received referral and add a thank-you amount. There is no Log Referral button here.",
        20,
        16,
    )
    card(
        s,
        Inches(8.85),
        Inches(1.55),
        Inches(3.9),
        Inches(4.9),
        "History",
        "Named list of your referrals, thank-yous, and one-to-ones so you can look back later.",
        20,
        16,
    )

    # 4 Give a referral
    s = prs.slides.add_slide(blank)
    header_bar(s, "Give a referral", "Always start from the dashboard — not My Referrals")
    footer(s, 4)
    add_bullets(
        s,
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5.3),
        [
            "On Dashboard, open My Activity and tap the Referrals card.",
            "To — search and select the chapter member you are referring to.",
            "Referral type — Inside or Outside.",
            "Optional status — Told them you would call, Given your card.",
            "Referral name (required), telephone, email, address, comments.",
            "Heat bar — how warm the lead is.",
            "Confirm. The slip is saved as Open for the receiver.",
        ],
        18,
        INK,
        10,
    )

    # 5 Receiver sees
    s = prs.slides.add_slide(blank)
    header_bar(s, "What the receiver sees", "The same slip, on their My Referrals list")
    footer(s, 5)
    add_bullets(
        s,
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5.3),
        [
            "The member you selected opens My Referrals.",
            "The slip appears on Received with status Open.",
            "It also stays on your Given tab — including after they convert it.",
            "Both of you are looking at the same referral, not two separate copies.",
        ],
        20,
        INK,
        14,
    )

    # 6 Convert
    s = prs.slides.add_slide(blank)
    header_bar(s, "Convert and log a thank you", "Only the receiver marks the referral converted")
    footer(s, 6)
    add_bullets(
        s,
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5.3),
        [
            "On My Referrals, open the Received tab.",
            "On an Open slip, click Mark as converted.",
            "A thank-you popup opens. Enter amount in rupees (₹).",
            "Choose business type: New or Repeat.",
            "Choose referral type: Inside, Outside, or Tier3+.",
            "Add comments if you want, then save.",
            "The slip shows as Converted. It still appears on Received.",
        ],
        18,
        INK,
        10,
    )

    # 7 Dashboard TYFCB
    s = prs.slides.add_slide(blank)
    header_bar(s, "Thank you from the dashboard", "TYFCBs on My Activity is the same thank-you")
    footer(s, 7)
    add_bullets(
        s,
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5.3),
        [
            "On Dashboard, open My Activity and tap TYFCBs.",
            "Thank you to — the member you are thanking.",
            "Amount (₹).",
            "Business type — New or Repeat.",
            "Referral type — Inside, Outside, or Tier3+.",
            "Comments, then Confirm.",
            "Use this for a thank-you that is not tied to converting a specific slip, or in addition to convert.",
        ],
        18,
        INK,
        10,
    )

    # 8 Who gets rupees
    s = prs.slides.add_slide(blank)
    header_bar(s, "Who gets the rupees", "Thank-you credit goes to one person only")
    footer(s, 8)
    card(
        s,
        Inches(0.55),
        Inches(1.55),
        Inches(6.0),
        Inches(4.9),
        "Credited",
        "The member who received the business — the thank-you recipient — gets the ₹ amount on their TYFCB total.",
        22,
        18,
    )
    card(
        s,
        Inches(6.8),
        Inches(1.55),
        Inches(6.0),
        Inches(4.9),
        "Not credited",
        "The giver does not get those rupees added to their TYFCB. They still keep the referral on Given, including after it is converted.",
        22,
        18,
    )

    # 9 One-to-one
    s = prs.slides.add_slide(blank)
    header_bar(s, "Log a one-to-one", "My Activity → One-to-Ones")
    footer(s, 9)
    add_bullets(
        s,
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5.3),
        [
            "On Dashboard, open My Activity and tap One-to-Ones.",
            "With — search and select the member you met.",
            "Initiated by — Me, Them, or Mutual.",
            "Where did you meet? — café, office, online, and so on.",
            "Date of the meeting.",
            "Topics of conversation.",
            "Confirm. You need a member in With before you can save.",
        ],
        18,
        INK,
        10,
    )

    # 10 Tabs
    s = prs.slides.add_slide(blank)
    header_bar(s, "My Referrals tabs", "Given, Received, and Converted")
    footer(s, 10)
    card(
        s,
        Inches(0.45),
        Inches(1.55),
        Inches(4.05),
        Inches(5.0),
        "Given",
        "Every referral you sent, including ones the other member has already converted.",
        20,
        16,
    )
    card(
        s,
        Inches(4.65),
        Inches(1.55),
        Inches(4.05),
        Inches(5.0),
        "Received",
        "Every referral sent to you, including converted ones. Convert from here.",
        20,
        16,
    )
    card(
        s,
        Inches(8.85),
        Inches(1.55),
        Inches(4.0),
        Inches(5.0),
        "Converted",
        "Referrals you received and marked converted. These also stay on Received.",
        20,
        16,
    )

    # 11 History
    s = prs.slides.add_slide(blank)
    header_bar(s, "History", "A named list of your slips")
    footer(s, 11)
    add_bullets(
        s,
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(5.3),
        [
            "Open History from the member sidebar.",
            "You will see named records for referrals, thank-yous, and one-to-ones.",
            "Use it to check who you referred, who you thanked, and who you met.",
            "Older PALMS counts from before named tracking may appear as “Before tracking”.",
        ],
        20,
        INK,
        14,
    )

    # 12 Help
    s = prs.slides.add_slide(blank)
    header_bar(s, "Need help?", "Recap of the three actions")
    footer(s, 12)
    add_bullets(
        s,
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(2.4),
        [
            "Login problems — contact your chapter admin. Access is invitation only.",
            "Give a referral — Dashboard → My Activity → Referrals.",
            "Thank you — convert on My Referrals, or Dashboard → My Activity → TYFCBs.",
            "One-to-one — Dashboard → My Activity → One-to-Ones.",
        ],
        18,
        INK,
        10,
    )
    add_round(s, Inches(0.55), Inches(4.2), Inches(12.2), Inches(2.5), PALE)
    add_textbox(
        s,
        Inches(0.85),
        Inches(4.45),
        Inches(11.6),
        Inches(2.0),
        "Remember: log new referrals from the dashboard. My Referrals is for tracking and converting. Thank-you rupees credit only the member who received the business.",
        18,
        True,
        NAVY,
    )

    out = Path(__file__).resolve().parent / "YEF-Member-Guide-Referrals-ThankYou-OneToOne.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
